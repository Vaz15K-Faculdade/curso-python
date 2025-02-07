from openpyxl import load_workbook

# Abrir arquivo excel
arquivo_excel = load_workbook('Animais.xlsx')

# Mostra o nome das Planilhas Existentes
print(arquivo_excel.sheetnames)

# Planilha 'animais'
planilha = arquivo_excel['animais']

# Metodo para printar direto
print(planilha.cell(2, 1).value)

total_linhas = planilha.max_row

# Percorre todas as linhas da planilha,
# Atribui a coluna 1 como nome
# Coluna 2 coloca como idade
for linha in range(2, total_linhas + 1):
    nome = planilha.cell(linha, 1).value
    idade = planilha.cell(linha, 2).value

    print(f'{nome} - {idade} anos')

# Atribui para a celula 11,1 os que estava antes so que Tudo maisculo
planilha.cell(11, 1).value = planilha.cell(11, 1).value.upper()

arquivo_excel.save()