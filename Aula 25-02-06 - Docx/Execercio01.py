"""
Faça um programa que preenche uma planilha
'carros.xlsx' com o dados abaixo:
"""

from openpyxl import Workbook

# Placa, Marca, Modelo, Cor
carros = [
    ["Placa", "Marca", "Modelo", "Cor"],
    ["AVX-0102", "Fiat", "Toro", "Branco"],
    ["TRF-2398", "Nissan", "Leaf", "Prata"],
    ["RWD-3891", "Honda", "City", "Preto"]
]

arquivo_excel = Workbook()

planilha = arquivo_excel.active

planilha.title = "Modelos"

for linha in range(len(carros)):
    for coluna in range(len(carros[linha])):
        planilha.cell(linha+1, coluna+1).value = carros[linha][coluna]
        
        if coluna == 1:
            planilha.cell(linha+1, coluna+1).value = planilha.cell(linha+1, coluna+1).value.upper()

arquivo_excel.save("carros.xlsx")