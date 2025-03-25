from smtplib import SMTP_SSL
from openpyxl import load_workbook
from email.mime.text import MIMEText
import json

def carregar_config():
    with open("config.json", "r") as arq:
        config = json.load(arq)

    return config

def enviar_email(destinatario, remetente, assunto, corpo, conexao):
    msg = MIMEText(corpo)
    msg["Subject"] = assunto
    msg["from"] = remetente

    conexao.sendmail(remetente, destinatario, msg.as_string())

def conectar_email(usuario, senha, servidor, porta):
    conexao = SMTP_SSL(servidor, porta)
    conexao.login(usuario, senha)

    return conexao

config = carregar_config()
conexao = conectar_email(config["usuario"], config["senha"], config["servidor"], config["porta"])
remetente = config["usuario"]

excel = load_workbook("inscricoes.xlsx")
planilha = excel.active

total_linhas = planilha.max_row

for linha in range(2, total_linhas + 1):
    num = planilha.cell(linha, 1).value
    nome = planilha.cell(linha, 2).value
    email = planilha.cell(linha, 3).value

    assunto = "Confirmação de inscrição"
    corpo = f"""
    Ola {nome}  
    Sua inscrição foi deferida
    Segue o seu numero de matricula {num}
    """

    enviar_email(email, remetente, assunto, corpo, conexao)